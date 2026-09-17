"""Session reports and lossless spectrum exports, independent of the GUI.

PDF and PowerPoint share a paginated page model. Tables/text remain editable;
only scientific plots/images are rasterized. Analysis arrays are never modified.
"""
from __future__ import annotations

import csv
from dataclasses import dataclass
from functools import lru_cache, wraps
from io import BytesIO
import os
from pathlib import Path
import re
from threading import RLock
import unicodedata

import numpy as np

PAGE_WIDTH, PAGE_HEIGHT = 960.0, 540.0
MARGIN, TABLE_FONT_SIZE, LINE_HEIGHT = 32.0, 10.0, 13.0
PROFILE_BATCH_SIZE = 6
# Layer-page layout: image (top) and spectrum (below) stacked in the left
# column; the parameter table runs tall down the right column and is shrunk to
# fit on a single page.
LEFT_X, LEFT_W = MARGIN, 464.0
IMAGE_TOP, IMAGE_H = 84.0, 250.0
PROFILE_TOP, PROFILE_H = 350.0, 162.0
RIGHT_X = 512.0
RIGHT_TABLE_TOP = 84.0
RIGHT_TABLE_AVAILABLE = PAGE_HEIGHT - RIGHT_TABLE_TOP - 26.0
TABLE_MIN_FONT = 5.0
ROI_COLOR = "#00d0ff"
_RENDER_LOCK = RLock()


def _serialized(function):
    @wraps(function)
    def wrapped(*args, **kwargs):
        with _RENDER_LOCK:
            return function(*args, **kwargs)
    return wrapped


def _safe_name(value: str) -> str:
    """Keep measurement/label names in their original language."""
    value = unicodedata.normalize("NFC", str(value))
    value = re.sub(r'[\\/:*?"<>|\x00-\x1f]+', "_", value)
    return re.sub(r"\s+", "_", value).strip(" ._")[:100] or "data"


def _unique_name(value: str, used: set[str]) -> str:
    base = _safe_name(value)
    candidate, number = base, 2
    while candidate.casefold() in used:
        candidate = f"{base}_{number}"
        number += 1
    used.add(candidate.casefold())
    return candidate


def _display_limits(data: np.ndarray, display: dict | None = None) -> tuple[float, float]:
    levels = (display or {}).get("levels")
    if levels is not None:
        limits = np.asarray(levels, dtype=float).reshape(-1)
        if limits.size == 2 and np.isfinite(limits).all() and limits[1] > limits[0]:
            return float(limits[0]), float(limits[1])
    finite = data[np.isfinite(data)]
    if not finite.size:
        return 0.0, 1.0
    lo, hi = np.percentile(finite, (1, 99))
    return float(lo), float(hi if hi > lo else lo + 1.0)


def _palette(name: str):
    """Match the viewer's explicit Jet/Fire control points exactly."""
    from matplotlib.colors import LinearSegmentedColormap

    if name.lower() == "jet":
        positions = [0, .125, .375, .625, .875, 1]
        colors = [(0, 0, 128), (0, 0, 255), (0, 255, 255),
                  (255, 255, 0), (255, 0, 0), (128, 0, 0)]
    elif name.lower() == "fire":
        positions = [0, .30, .55, .78, 1]
        colors = [(0, 0, 0), (150, 0, 0), (255, 60, 0), (255, 220, 0), (255, 255, 255)]
    else:
        positions, colors = [0, 1], [(0, 0, 0), (255, 255, 255)]
    return LinearSegmentedColormap.from_list(
        name, list(zip(positions, np.asarray(colors, dtype=float) / 255)), N=256
    )


def _image_array(array, display: dict | None = None) -> np.ndarray:
    data = np.asarray(array)
    if data.ndim == 3 and data.shape[-1] in (3, 4):
        if data.dtype == np.uint8:
            return data.copy()
        maximum = 255.0 if np.issubdtype(data.dtype, np.integer) else 1.0
        data = np.nan_to_num(data.astype(float), nan=0.0, posinf=maximum, neginf=0.0)
        return np.rint(np.clip(data / maximum, 0, 1) * 255).astype(np.uint8)
    if data.ndim != 2:
        raise ValueError(f"Report layer image must be 2D or RGB/RGBA, received {data.shape}.")
    data = data.astype(float)
    lo, hi = _display_limits(data, display)
    normalized = np.clip((np.nan_to_num(data, nan=lo, posinf=hi, neginf=lo) - lo) / (hi - lo), 0, 1)
    palette = (display or {}).get("palette", "Grayscale")
    if str(palette).lower() in ("gray", "grey", "grayscale"):
        return np.rint(normalized * 255).astype(np.uint8)
    return np.rint(_palette(str(palette))(normalized)[..., :3] * 255).astype(np.uint8)


@_serialized
def export_layer_images(layers: list[dict], directory: Path) -> list[Path]:
    """Export one display-faithful PNG per layer, without name collisions."""
    from PIL import Image

    directory = Path(directory)
    directory.mkdir(parents=True, exist_ok=True)
    outputs, used = [], set()
    for layer in layers:
        if layer.get("image") is None:
            continue
        path = directory / f"{_unique_name(layer.get('name', 'layer'), used)}.png"
        Image.fromarray(_image_array(layer["image"], layer.get("display"))).save(path)
        outputs.append(path)
    return outputs


@_serialized
def export_image_set(image_set: dict, directory: Path) -> list[Path]:
    """Export one image type into its own folder.

    A ``stack`` set writes one display-faithful PNG per energy frame (named by
    index and energy); a ``map`` set writes the single map image.  ``directory``
    is the parent; a sub-folder named after the set is created inside it.
    """
    from PIL import Image

    name = str(image_set.get("name", "images"))
    folder = Path(directory) / _safe_name(name)
    folder.mkdir(parents=True, exist_ok=True)
    display = image_set.get("display")
    data = image_set.get("data")
    outputs = []
    if data is None:
        return outputs
    if image_set.get("kind") == "segmentation":
        # ``data`` is a list of saved segmentation groups, each with a combined
        # all-clusters colour image and one image per cluster.  With several
        # saved labels, each goes in its own sub-folder.
        groups = list(data)
        for group in groups:
            label = str(group.get("label", "segmentation"))
            target = folder / _safe_name(label) if len(groups) > 1 else folder
            target.mkdir(parents=True, exist_ok=True)
            combined = group.get("combined")
            if combined is not None:
                path = target / "all_clusters.png"
                Image.fromarray(_image_array(combined)).save(path)
                outputs.append(path)
            for cluster in group.get("clusters", []):
                cid = cluster.get("id")
                cimg = cluster.get("image")
                if cimg is None:
                    continue
                path = target / f"cluster_{int(cid):03d}.png"
                Image.fromarray(_image_array(cimg)).save(path)
                outputs.append(path)
        return outputs
    if image_set.get("kind") == "stack":
        stack = np.asarray(data)
        energies = image_set.get("energies")
        energies = np.asarray(energies) if energies is not None else None
        # One scaling for the whole stack so contrast is identical across every
        # energy frame (comparable frame to frame).  Use the stack-wide finite
        # min–max unless explicit levels are supplied.  Per-frame scaling would
        # both change contrast between frames and blow out low-signal frames
        # (e.g. OD − pre-edge, where signal is sparse over a near-zero background).
        if display and display.get("levels") is not None:
            levels = tuple(display["levels"])
        else:
            finite = stack[np.isfinite(stack)]
            lo = float(finite.min()) if finite.size else 0.0
            hi = float(finite.max()) if finite.size else 1.0
            levels = (lo, hi if hi > lo else lo + 1.0)
        frame_display = {**(display or {}), "levels": levels}
        for i in range(len(stack)):
            suffix = ""
            if energies is not None and i < energies.size and np.isfinite(energies[i]):
                suffix = f"_{float(energies[i]):.2f}eV"
            path = folder / f"{_safe_name(name)}_{i:03d}{suffix}.png"
            Image.fromarray(_image_array(stack[i], frame_display)).save(path)
            outputs.append(path)
    else:
        path = folder / f"{_safe_name(name)}.png"
        Image.fromarray(_image_array(data, display)).save(path)
        outputs.append(path)
    return outputs


def _validated_spectra(spectra: dict[str, object], energies) -> tuple[np.ndarray, list[np.ndarray]]:
    axis = np.asarray(energies, dtype=float)
    if axis.ndim != 1:
        raise ValueError("The energy axis must be one-dimensional.")
    arrays = []
    for name, values in spectra.items():
        values = np.asarray(values, dtype=float)
        if values.ndim != 1 or values.size != axis.size:
            raise ValueError(f"Spectrum '{name}' has shape {values.shape}; expected {axis.size} energy values.")
        arrays.append(values)
    return axis, arrays


def export_spectra_csv(spectra: dict[str, object], energies: np.ndarray, path: Path) -> Path:
    """Preserve full float precision and every row; reject mismatched axes."""
    axis, arrays = _validated_spectra(spectra, energies)
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.writer(handle)
        writer.writerow(["energy_eV", *spectra])
        writer.writerows(zip(axis, *arrays, strict=True))
    return path


def export_spectra_bundle(spectra: dict[str, object], energies: np.ndarray,
                          mapping_spectra: dict[str, dict], directory: Path,
                          stem: str = "report") -> list[Path]:
    """Write aggregate CSV and separate CSVs containing each map's used spectra.

    Map entries are ``{name: values}`` or ``{'spectra': {...}, 'energies': axis}``.
    Callers supply only actually used spectra. Validate the complete bundle
    before opening any file, so mismatched lengths never silently truncate rows.
    """
    entries = [(f"{stem}_all_spectra", spectra, energies)]
    for map_name, item in mapping_spectra.items():
        if "spectra" in item and isinstance(item["spectra"], dict):
            values, axis = item["spectra"], item.get("energies", energies)
        else:
            values, axis = item, energies
        if values:
            entries.append((f"{stem}_{map_name}_spectra", values, axis))
    for _, values, axis in entries:
        _validated_spectra(values, axis)
    used, outputs = set(), []
    for name, values, axis in entries:
        outputs.append(export_spectra_csv(values, axis, Path(directory) / f"{_unique_name(name, used)}.csv"))
    return outputs


@lru_cache(maxsize=1)
def _report_font() -> tuple[str, str | None, str]:
    """Embed a system TrueType font; macOS Arial Unicode includes Japanese."""
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont

    candidates = [os.environ.get("MUNEXAFS_REPORT_FONT", ""),
                  os.environ.get("MUAXIS_REPORT_FONT", ""),  # legacy setting
                  "/System/Library/Fonts/Supplemental/Arial Unicode.ttf",
                  "/Library/Fonts/Arial Unicode.ttf", "C:/Windows/Fonts/msgothic.ttc",
                  "/usr/share/fonts/truetype/fonts-japanese-gothic.ttf",
                  "/usr/share/fonts/opentype/ipafont-gothic/ipag.ttf",
                  "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"]
    for candidate in candidates:
        if not candidate or not Path(candidate).is_file():
            continue
        try:
            font = TTFont("muNEXAFSReport", candidate)
            pdfmetrics.registerFont(font)
            family = font.face.familyName
            if isinstance(family, bytes):
                family = family.decode("utf-8", errors="replace")
            return "muNEXAFSReport", candidate, str(family)
        except Exception:
            continue
    return "Helvetica", None, "Arial"


def _string(value) -> str:
    if value is None:
        return "Not recorded"
    if isinstance(value, (bool, np.bool_)):
        return "Yes" if value else "No"
    if isinstance(value, np.ndarray):
        value = value.tolist()
    if isinstance(value, dict):
        return "; ".join(f"{key}: {_string(item)}" for key, item in value.items())
    if isinstance(value, (list, tuple)):
        return ", ".join(_string(item) for item in value)
    return str(value)


def _wrap(value: str, width: float, size: float = TABLE_FONT_SIZE) -> list[str]:
    """Hard-wrap even long paths or Japanese labels, without dropping text."""
    from reportlab.pdfbase.pdfmetrics import stringWidth

    font, result = _report_font()[0], []
    for paragraph in str(value).split("\n"):
        line = ""
        for char in paragraph:
            if line and stringWidth(line + char, font, size) > width:
                result.append(line)
                line = ""
            line += char
        result.append(line)
    return result or [""]


@dataclass
class ReportRow:
    cells: list[str]
    height: float


@dataclass
class ReportPage:
    title: str
    subtitle: str
    columns: list[str]
    widths: list[float]
    rows: list[ReportRow]
    table_top: float
    layer_plot: bytes | None = None
    profile_plot: bytes | None = None
    image_notice: str = ""
    profile_notice: str = ""
    layout: str = "overview"          # "overview" (full-width) or "layer" (right column)
    table_x: float = MARGIN
    table_font: float = TABLE_FONT_SIZE
    header_height: float = 24.0


def _paginate_rows(rows: list[list[str]], widths: list[float], available: float) -> list[list[ReportRow]]:
    pages, current, used = [], [], 0.0
    max_lines = max(1, int((available - 12) // LINE_HEIGHT))
    for raw in rows:
        wrapped = [_wrap(cell, width - 16) for cell, width in zip(raw, widths, strict=True)]
        count = max(map(len, wrapped))
        for start in range(0, count, max_lines):
            lines = [cell[start:start + max_lines] for cell in wrapped]
            height = max(map(len, lines)) * LINE_HEIGHT + 12
            if current and used + height > available:
                pages.append(current)
                current, used = [], 0.0
            current.append(ReportRow(["\n".join(cell) for cell in lines], height))
            used += height
    if current or not pages:
        pages.append(current)
    return pages


def _table_rows_at(rows: list[list[str]], widths: list[float], font: float) -> list[ReportRow]:
    """Wrap each cell at ``font`` and return rows with their pixel heights."""
    line_h = font + 3.0
    built = []
    for raw in rows:
        wrapped = [_wrap(cell, width - 16, font) for cell, width in zip(raw, widths, strict=True)]
        count = max(map(len, wrapped)) or 1
        built.append(ReportRow(["\n".join(cell) for cell in wrapped], count * line_h + 8.0))
    return built


def _fit_table(rows: list[list[str]], widths: list[float], available: float,
               max_font: float = TABLE_FONT_SIZE, min_font: float = TABLE_MIN_FONT):
    """Largest font (0.5 pt steps) at which the whole table fits ``available``
    height on one page; falls back to ``min_font`` if nothing fits."""
    font = max_font
    while font >= min_font:
        built = _table_rows_at(rows, widths, font)
        header = font + 14.0
        if header + sum(row.height for row in built) <= available:
            return font, built, header
        font -= 0.5
    # Beyond what even the smallest font can fit: keep the rows that fit on the
    # single page and replace the remainder with a pointer to the full listing.
    built = _table_rows_at(rows, widths, min_font)
    header = min_font + 14.0
    kept, used = [], header
    for i, row in enumerate(built):
        if not kept or used + row.height <= available:
            kept.append(row); used += row.height
        else:
            more = _table_rows_at(
                [["…", f"{len(built) - i} more parameter(s) omitted here; see the "
                        "parameter overview page and the exported CSV"]], widths, min_font)[0]
            if used + more.height <= available:
                kept.append(more)
            break
    return min_font, kept, header


def _figure_png(figure) -> bytes:
    from matplotlib.backends.backend_agg import FigureCanvasAgg

    buffer = BytesIO()
    FigureCanvasAgg(figure).print_png(buffer)
    return buffer.getvalue()


def _plot_font(size=10):
    from matplotlib.font_manager import FontProperties

    path = _report_font()[1]
    return FontProperties(fname=path, size=size) if path else FontProperties(size=size)


def _layer_plot(layer: dict) -> bytes | None:
    from matplotlib.figure import Figure
    from matplotlib.patches import Rectangle

    if layer.get("image") is None:
        return None
    data = np.asarray(layer["image"])
    # A squarer figure with tight margins so the map fills the panel (larger).
    figure = Figure(figsize=(4.9, 3.9), dpi=185, facecolor="white")
    axis = figure.add_subplot(111)
    figure.subplots_adjust(left=.13, right=.86, bottom=.12, top=.96)
    if data.ndim == 3:
        axis.imshow(_image_array(data), origin="upper", interpolation="nearest", aspect="equal")
    else:
        lo, hi = _display_limits(data, layer.get("display"))
        palette = (layer.get("display") or {}).get("palette", "Grayscale")
        image = axis.imshow(np.ma.masked_invalid(data), origin="upper", interpolation="nearest",
                            aspect="equal", cmap=_palette(str(palette)), vmin=lo, vmax=hi)
        colorbar = figure.colorbar(image, ax=axis, pad=.03, fraction=.045)
        colorbar.ax.tick_params(labelsize=9)
        unit = (layer.get("display") or {}).get("units")
        if unit:
            colorbar.set_label(str(unit), fontproperties=_plot_font(9))
    for number, roi in enumerate(layer.get("rois") or [], 1):
        x0, y0, x1, y1 = roi
        axis.add_patch(Rectangle((x0 - 0.5, y0 - 0.5), x1 - x0, y1 - y0,
                                 fill=False, edgecolor=ROI_COLOR, linewidth=1.6))
        axis.text(x0, y0 - 1, f"ROI {number}", color=ROI_COLOR, va="bottom", ha="left",
                  fontproperties=_plot_font(8))
    axis.set_xlabel("X (pixel)", fontproperties=_plot_font())
    axis.set_ylabel("Y (pixel)", fontproperties=_plot_font())
    axis.tick_params(labelsize=9)
    return _figure_png(figure)


def _profile_plot(profiles: list[dict], energies, ylabel: str) -> bytes | None:
    from matplotlib.figure import Figure
    from matplotlib.colors import to_rgba

    if not profiles:
        return None
    figure = Figure(figsize=(7.0, 3.9), dpi=170, facecolor="white")
    axis = figure.add_subplot(111)
    figure.subplots_adjust(left=.13, right=.97, bottom=.34, top=.94)
    for profile in profiles:
        values = np.asarray(profile["values"], dtype=float)
        energy = profile.get("energies", energies)
        if energy is None:
            energy = np.arange(values.size)
        energy, arrays = _validated_spectra({profile.get("name", "Spectrum"): values}, energy)
        color = profile.get("color")
        if isinstance(color, (tuple, list, np.ndarray)):
            color = np.asarray(color, dtype=float)
            if np.max(color) > 1:
                color = color / 255
        try:
            color = to_rgba(color) if color is not None else None
        except (TypeError, ValueError):
            color = None
        label = "\n".join(_wrap(str(profile.get("name", "Spectrum")), 190, 8))
        axis.plot(energy, arrays[0], lw=1.4, label=label, color=color)
    axis.set_xlabel("Energy (eV)" if energies is not None else "Energy index (axis unavailable)",
                    fontproperties=_plot_font())
    axis.set_ylabel(ylabel, fontproperties=_plot_font())
    axis.tick_params(labelsize=9)
    axis.grid(alpha=.22, linewidth=.5)
    axis.legend(loc="upper center", bbox_to_anchor=(.5, -.26), ncol=2,
                frameon=False, prop=_plot_font(8), borderaxespad=0)
    return _figure_png(figure)


@_serialized
def build_report_pages(layers: list[dict], metadata: dict, *, energies=None) -> list[ReportPage]:
    """Build complete overview/step pages used by both document exporters."""
    if energies is None:
        energies = metadata.get("energies")
    overview = []
    measurement = metadata.get("measurement_parameters") or {
        key: value for key, value in metadata.items()
        if key not in ("analysis_parameters", "energies", "measurement_parameters")
    }
    for key, value in measurement.items():
        overview.append(["Measurement", str(key), _string(value)])
    analysis = metadata.get("analysis_parameters") or {
        layer.get("name", f"Step {i + 1}"): layer.get("parameters", {}) for i, layer in enumerate(layers)
    }
    for step, parameters in analysis.items():
        if not isinstance(parameters, dict):
            parameters = {"Conditions": parameters}
        if not parameters:
            parameters = {"Conditions": "Not recorded"}
        for key, value in parameters.items():
            overview.append([str(step), str(key), _string(value)])
    if not overview:
        overview = [["Measurement", "Analysis parameters", "No parameters recorded"]]
    widths = [185., 235., PAGE_WIDTH - MARGIN * 2 - 420.]
    groups = _paginate_rows(overview, widths, 366)
    pages, subtitle = [], str(metadata.get("title", "STXM / NEXAFS analysis report"))
    for i, rows in enumerate(groups):
        title = "Analysis parameter overview" + (f" ({i + 1}/{len(groups)})" if len(groups) > 1 else "")
        pages.append(ReportPage(title, subtitle, ["Step / group", "Parameter", "Value"], widths, rows, 108))
    for index, layer in enumerate(layers, 1):
        rows = []
        if layer.get("source"):
            rows.append(["Input layer", _string(layer["source"])])
        if layer.get("condition"):
            rows.append(["Summary", _string(layer["condition"])])
        rows += [[str(key), _string(value)] for key, value in (layer.get("parameters") or {}).items()]
        if layer.get("display"):
            rows.append(["Image display", _string(layer["display"])])
        profiles = list(layer.get("profiles") or [])
        if profiles:
            rows.append(["Profiles included", ", ".join(
                str(profile.get("name", "Spectrum")) for profile in profiles
            )])
        if not rows:
            rows = [["Conditions", "No analysis conditions recorded"]]
        # Tall table down the right column: one page per layer, font shrunk to fit.
        condition_width = 150.0
        value_width = (PAGE_WIDTH - MARGIN - RIGHT_X) - condition_width
        widths = [condition_width, value_width]
        font, built, header_height = _fit_table(rows, widths, RIGHT_TABLE_AVAILABLE)
        layer_image = _layer_plot(layer)
        profile_image = _profile_plot(profiles, energies, layer.get("profile_ylabel", "Optical density (OD)"))
        name = f"{index:02d}  {layer.get('name', 'Analysis step')}"
        pages.append(ReportPage(
            name, subtitle, ["Analysis condition", "Value"], widths, built, RIGHT_TABLE_TOP,
            layer_image, profile_image,
            "No layer image saved" if layer_image is None else "",
            "No profile saved for this step" if not profiles else "",
            layout="layer", table_x=RIGHT_X, table_font=font, header_height=header_height))
    return pages


def _pdf_text(pdf, text: str, x: float, top: float, size: float = 10,
              width: float | None = None, color: str = "#172B3A"):
    from reportlab.lib.colors import HexColor

    pdf.setFillColor(HexColor(color))
    pdf.setFont(_report_font()[0], size)
    lines = _wrap(text, width, size) if width else str(text).split("\n")
    for i, line in enumerate(lines):
        pdf.drawString(x, PAGE_HEIGHT - top - size - i * (size + 3), line)


def _pdf_table(pdf, page: ReportPage):
    from reportlab.lib.colors import HexColor

    x0, font, header_h, width = page.table_x, page.table_font, page.header_height, sum(page.widths)
    y = page.table_top
    pdf.setFillColor(HexColor("#E6EEF3"))
    pdf.rect(x0, PAGE_HEIGHT - y - header_h, width, header_h, stroke=0, fill=1)
    x = x0
    for text, column_width in zip(page.columns, page.widths, strict=True):
        _pdf_text(pdf, text, x + 8, y + (header_h - font) / 2, font)
        x += column_width
    y += header_h
    for number, row in enumerate(page.rows):
        pdf.setFillColor(HexColor("#F5F8FA" if number % 2 == 0 else "#FFFFFF"))
        pdf.rect(x0, PAGE_HEIGHT - y - row.height, width, row.height, stroke=0, fill=1)
        x = x0
        for text, column_width in zip(row.cells, page.widths, strict=True):
            _pdf_text(pdf, text, x + 8, y + 5, font)
            x += column_width
        pdf.setStrokeColor(HexColor("#D6E0E7"))
        pdf.line(x0, PAGE_HEIGHT - y - row.height, x0 + width, PAGE_HEIGHT - y - row.height)
        y += row.height


@_serialized
def export_pdf(layers: list[dict], metadata: dict, directory: Path, filename: str, *, energies=None) -> Path:
    from reportlab.lib.utils import ImageReader
    from reportlab.pdfgen import canvas

    pages = build_report_pages(layers, metadata, energies=energies)
    path = Path(directory) / filename
    path.parent.mkdir(parents=True, exist_ok=True)
    pdf = canvas.Canvas(str(path), pagesize=(PAGE_WIDTH, PAGE_HEIGHT))
    pdf.setTitle(str(metadata.get("title", "muNEXAFS analysis report")))
    pdf.setAuthor("muNEXAFS")
    for index, page in enumerate(pages, 1):
        _pdf_text(pdf, page.title, MARGIN, 19, 20, PAGE_WIDTH - 2 * MARGIN)
        _pdf_text(pdf, page.subtitle, MARGIN, 55, 10, PAGE_WIDTH - 2 * MARGIN, "#526B7A")
        if page.layout == "layer":
            # Left column: layer image on top, spectrum below (stacked).
            _pdf_text(pdf, "Layer image", LEFT_X, IMAGE_TOP - 14, 11)
            if page.layer_plot:
                pdf.drawImage(ImageReader(BytesIO(page.layer_plot)), LEFT_X, PAGE_HEIGHT - IMAGE_TOP - IMAGE_H,
                              width=LEFT_W, height=IMAGE_H, preserveAspectRatio=True, anchor="c")
            else:
                _pdf_text(pdf, page.image_notice, LEFT_X + 15, IMAGE_TOP + IMAGE_H / 2, 12, 400, "#637381")
            _pdf_text(pdf, "Profiles", LEFT_X, PROFILE_TOP - 14, 11)
            if page.profile_plot:
                pdf.drawImage(ImageReader(BytesIO(page.profile_plot)), LEFT_X, PAGE_HEIGHT - PROFILE_TOP - PROFILE_H,
                              width=LEFT_W, height=PROFILE_H, preserveAspectRatio=True, anchor="c")
            else:
                _pdf_text(pdf, page.profile_notice, LEFT_X + 15, PROFILE_TOP + PROFILE_H / 2, 12, 400, "#637381")
        _pdf_table(pdf, page)
        _pdf_text(pdf, "muNEXAFS | STXM–NEXAFS analysis", MARGIN, 518, 8, color="#637381")
        _pdf_text(pdf, f"{index} / {len(pages)}", PAGE_WIDTH - 80, 518, 8, color="#637381")
        pdf.showPage()
    pdf.save()
    return path


def _pptx_text(slide, text: str, x: float, top: float, width: float, height: float,
               size=10, color="172B3A"):
    from pptx.dml.color import RGBColor
    from pptx.util import Pt

    frame = slide.shapes.add_textbox(Pt(x), Pt(top), Pt(width), Pt(height)).text_frame
    frame.clear()
    frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = 0
    frame.word_wrap = True
    for index, line in enumerate(str(text).split("\n")):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.text = line
        paragraph.font.name = _report_font()[2]
        paragraph.font.size = Pt(size)
        paragraph.font.color.rgb = RGBColor.from_string(color)
        paragraph.space_before = paragraph.space_after = Pt(0)
    return frame


def _pptx_table(slide, page: ReportPage):
    from pptx.dml.color import RGBColor
    from pptx.util import Pt
    from pptx.enum.text import MSO_ANCHOR

    font, header_h = page.table_font, page.header_height
    table = slide.shapes.add_table(len(page.rows) + 1, len(page.columns), Pt(page.table_x), Pt(page.table_top),
                                   Pt(sum(page.widths)), Pt(header_h + sum(row.height for row in page.rows))).table
    for column, width in zip(table.columns, page.widths, strict=True):
        column.width = Pt(width)
    table.rows[0].height = Pt(header_h)
    for index, row in enumerate(page.rows, 1):
        table.rows[index].height = Pt(row.height)
    for row_index, cells in enumerate([page.columns] + [row.cells for row in page.rows]):
        for col_index, text in enumerate(cells):
            cell = table.cell(row_index, col_index)
            cell.text = text
            cell.margin_left = cell.margin_right = Pt(8)
            cell.margin_top = cell.margin_bottom = Pt(3)
            cell.vertical_anchor = MSO_ANCHOR.TOP
            cell.fill.solid()
            cell.fill.fore_color.rgb = RGBColor.from_string("E6EEF3" if row_index == 0 else
                                                           ("F5F8FA" if row_index % 2 else "FFFFFF"))
            for paragraph in cell.text_frame.paragraphs:
                paragraph.font.name = _report_font()[2]
                paragraph.font.size = Pt(font)
                paragraph.font.bold = row_index == 0
                paragraph.font.color.rgb = RGBColor.from_string("172B3A")
                paragraph.line_spacing = Pt(font + 3.0)
                paragraph.space_after = paragraph.space_before = Pt(0)


@_serialized
def export_pptx(layers: list[dict], metadata: dict, directory: Path, filename: str, *, energies=None) -> Path:
    from pptx import Presentation
    from pptx.util import Pt
    from PIL import Image

    pages = build_report_pages(layers, metadata, energies=energies)
    presentation = Presentation()
    presentation.slide_width, presentation.slide_height = Pt(PAGE_WIDTH), Pt(PAGE_HEIGHT)
    presentation.core_properties.title = str(metadata.get("title", "muNEXAFS analysis report"))
    for index, page in enumerate(pages, 1):
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        _pptx_text(slide, page.title, MARGIN, 19, PAGE_WIDTH - 2 * MARGIN, 32, 20)
        _pptx_text(slide, page.subtitle, MARGIN, 55, PAGE_WIDTH - 2 * MARGIN, 25, 10, "526B7A")
        if page.layout == "layer":
            for label, data, notice, top, box_h in (
                    ("Layer image", page.layer_plot, page.image_notice, IMAGE_TOP, IMAGE_H),
                    ("Profiles", page.profile_plot, page.profile_notice, PROFILE_TOP, PROFILE_H)):
                _pptx_text(slide, label, LEFT_X, top - 14, LEFT_W, 16, 11)
                if data:
                    with Image.open(BytesIO(data)) as image:
                        ratio = min(LEFT_W / image.width, box_h / image.height)
                        w, h = image.width * ratio, image.height * ratio
                    slide.shapes.add_picture(BytesIO(data), Pt(LEFT_X + (LEFT_W - w) / 2), Pt(top + (box_h - h) / 2),
                                             width=Pt(w), height=Pt(h))
                else:
                    _pptx_text(slide, notice, LEFT_X + 15, top + box_h / 2, 400, 40, 12, "637381")
        _pptx_table(slide, page)
        _pptx_text(slide, "muNEXAFS | STXM–NEXAFS analysis", MARGIN, 518, 500, 14, 8, "637381")
        _pptx_text(slide, f"{index} / {len(pages)}", PAGE_WIDTH - 80, 518, 55, 14, 8, "637381")
    path = Path(directory) / filename
    path.parent.mkdir(parents=True, exist_ok=True)
    presentation.save(path)
    return path
